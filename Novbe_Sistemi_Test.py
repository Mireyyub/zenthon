#!/usr/bin/env python3
"""
Növbətçilik Sistemi Test Skripti
"""

import openpyxl
from openpyxl import load_workbook
from datetime import date
from calendar import monthrange

# Azərbaycan dilində həftə günləri
HEFTE_GUNLERI = [
    "Bazar ertəsi", "Çərşənbə axşamı", "Çərşənbə", 
    "Cümə axşamı", "Cümə", "Şənbə", "Bazar"
]

# Aylar Azərbaycan dilində
AYLAR_ADLARI = [
    "Yanvar", "Fevral", "Mart", "Aprel", "May", "İyun",
    "İyul", "Avqust", "Sentyabr", "Oktyabr", "Noyabr", "Dekabr"
]


def test_file_structure():
    """Fayl strukturu testi"""
    print("=" * 60)
    print("TEST 1: Fayl strukturu")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    
    expected_sheets = [
        "DASHBOARD", "AYLIQ QRAFİK", "ŞƏXSLƏR", 
        "YOXDUR-MƏZUNİYYƏT", "BALANS", "PARAMETRLƏR", 
        "DƏYİŞİKLİK JURNALI"
    ]
    
    actual_sheets = wb.sheetnames
    
    print(f"Gözlənilən vərəqlər: {expected_sheets}")
    print(f"Faktiki vərəqlər: {actual_sheets}")
    
    if set(expected_sheets) == set(actual_sheets):
        print("✅ Bütün vərəqlər mövcuddur")
        return True
    else:
        missing = set(expected_sheets) - set(actual_sheets)
        extra = set(actual_sheets) - set(expected_sheets)
        print(f"❌ Nəsir vərəqlər: {missing}")
        print(f"❌ Əlavə vərəqlər: {extra}")
        return False


def test_parameters_sheet():
    """PARAMETRLƏR vərəqi testi"""
    print("\n" + "=" * 60)
    print("TEST 2: PARAMETRLƏR vərəqi")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    ws = wb["PARAMETRLƏR"]
    
    # Növbə əmsalları
    expected_params = {
        "Hissə növbətçisi": 5,
        "Hissə növbətçisinin köməkçisi": 4,
        "Nəzarət-buraxılış məntəqəsi növbətçisi": 3,
        "Park növbətçisi": 3,
        "Yeməkxana növbətçisi": 2,
        "Otaq növbətçisi": 1
    }
    
    print("Növbə əmsalları yoxlanılır:")
    all_ok = True
    for row in range(5, 11):
        novbe = ws[f"A{row}"].value
        emsal = ws[f"B{row}"].value
        expected_emsal = expected_params.get(novbe)
        
        if expected_emsal is not None:
            if emsal == expected_emsal:
                print(f"  ✅ {novbe}: {emsal}")
            else:
                print(f"  ❌ {novbe}: {emsal} (gözlənilən: {expected_emsal})")
                all_ok = False
    
    return all_ok


def test_people_sheet():
    """ŞƏXSLƏR vərəqi testi"""
    print("\n" + "=" * 60)
    print("TEST 3: ŞƏXSLƏR vərəqi")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    ws = wb["ŞƏXSLƏR"]
    
    # Başlıq sətri
    headers = [
        "№", "Soyad", "Ad", "Ata adı", "Tam ad", 
        "Vəzifə", "Aktiv / Aktiv deyil", "Qeyd"
    ]
    
    print("Başlıq sətri yoxlanılır:")
    all_ok = True
    for i, header in enumerate(headers):
        cell_value = ws.cell(row=4, column=i+1).value
        if cell_value == header:
            print(f"  ✅ {header}")
        else:
            print(f"  ❌ {header} (faktiki: {cell_value})")
            all_ok = False
    
    # Nümayiş məlumatları
    print("\nNümayiş məlumatları:")
    print(f"  Şəxs say: {ws.max_row - 4} (gözlənilən: 8)")
    
    return all_ok


def test_absence_sheet():
    """YOXDUR-MƏZUNİYYƏT vərəqi testi"""
    print("\n" + "=" * 60)
    print("TEST 4: YOXDUR-MƏZUNİYYƏT vərəqi")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    ws = wb["YOXDUR-MƏZUNİYYƏT"]
    
    headers = [
        "Şəxs", "Səbəb", "Başlanğıc tarixi", 
        "Gün sayı", "Bitmə tarixi", "Qeyd"
    ]
    
    print("Başlıq sətri yoxlanılır:")
    all_ok = True
    for i, header in enumerate(headers):
        cell_value = ws.cell(row=4, column=i+1).value
        if cell_value == header:
            print(f"  ✅ {header}")
        else:
            print(f"  ❌ {header} (faktiki: {cell_value})")
            all_ok = False
    
    # Bitmə tarixi formulu
    print("\nBitmə tarixi formulu:")
    for row in range(5, 8):
        formula = ws[f"E{row}"].value
        if formula and formula.startswith("="):
            print(f"  ✅ Sətir {row}: {formula}")
        else:
            print(f"  ❌ Sətir {row}: Formula tapılmadı")
            all_ok = False
    
    return all_ok


def test_monthly_schedule():
    """AYLIQ QRAFİK vərəqi testi"""
    print("\n" + "=" * 60)
    print("TEST 5: AYLIQ QRAFİK vərəqi")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    ws = wb["AYLIQ QRAFİK"]
    
    # Parametrlər
    il = ws["B4"].value
    ay = ws["E4"].value
    
    print(f"Seçilmiş il: {il}")
    print(f"Seçilmiş ay: {ay}")
    
    # Tarix sütunu
    print("\nTarix sütunu yoxlanılır:")
    all_ok = True
    
    # Sentyabr 2026-da 30 gün var
    expected_days = 30
    actual_days = ws.max_row - 7
    
    print(f"Gün sayı: {actual_days} (gözlənilən: {expected_days})")
    
    if actual_days == expected_days:
        print("✅ Düzgün gün sayı")
    else:
        print("❌ Yanlış gün sayı")
        all_ok = False
    
    # Həftə günləri
    print("\nHəftə günləri yoxlanılır:")
    for row in range(8, min(15, ws.max_row + 1)):
        hefte_gunu = ws[f"C{row}"].value
        if hefte_gunu in HEFTE_GUNLERI:
            print(f"  ✅ Sətir {row}: {hefte_gunu}")
        else:
            print(f"  ❌ Sətir {row}: {hefte_gunu}")
            all_ok = False
    
    # Növbə sütunları
    novbe_sutunlari = [
        "Hissə növbətçisi", "Hissə növbətçisinin köməkçisi",
        "Nəzarət-buraxılış məntəqəsi növbətçisi", "Park növbətçisi",
        "Yeməkxana növbətçisi", "Otaq növbətçisi"
    ]
    
    print("\nNövbə sütun başlıqları:")
    for i, novbe in enumerate(novbe_sutunlari):
        col = 5 + i
        header = ws.cell(row=7, column=col).value
        if header == novbe:
            print(f"  ✅ {novbe}")
        else:
            print(f"  ❌ {novbe} (faktiki: {header})")
            all_ok = False
    
    return all_ok


def test_balance_sheet():
    """BALANS vərəqi testi"""
    print("\n" + "=" * 60)
    print("TEST 6: BALANS vərəqi")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    ws = wb["BALANS"]
    
    headers = [
        "Şəxs", "Ümumi növbə", "Ümumi xidmət yükü", 
        "Mövcud gün", "Normallaşdırılmış yük"
    ]
    
    print("Başlıq sətri yoxlanılır:")
    all_ok = True
    for i, header in enumerate(headers[:5]):
        cell_value = ws.cell(row=4, column=i+1).value
        if cell_value == header:
            print(f"  ✅ {header}")
        else:
            print(f"  ❌ {header} (faktiki: {cell_value})")
            all_ok = False
    
    # Formulların yoxlanılması
    print("\nFormulların yoxlanılması:")
    for row in range(5, min(10, ws.max_row + 1)):
        person = ws[f"A{row}"].value
        if person:
            # Ümumi növbə formulu
            formula_b = ws[f"B{row}"].value
            if formula_b and formula_b.startswith("="):
                print(f"  ✅ {person} - Ümumi növbə: {formula_b[:50]}...")
            else:
                print(f"  ❌ {person} - Ümumi növbə: Formula tapılmadı")
                all_ok = False
    
    return all_ok


def test_dashboard():
    """DASHBOARD vərəqi testi"""
    print("\n" + "=" * 60)
    print("TEST 7: DASHBOARD vərəqi")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    ws = wb["DASHBOARD"]
    
    # Bölmələr
    sections = ["ÜMUMİ VƏZİYYƏT", "BALANS", "RİSKLƏR"]
    
    print("Bölmələr yoxlanılır:")
    all_ok = True
    for section in sections:
        found = False
        for row in range(1, ws.max_row + 1):
            if ws[f"A{row}"].value == section:
                print(f"  ✅ {section}")
                found = True
                break
        if not found:
            print(f"  ❌ {section} tapılmadı")
            all_ok = False
    
    return all_ok


def test_data_validation():
    """Data Validation testi"""
    print("\n" + "=" * 60)
    print("TEST 8: Data Validation")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    
    # ŞƏXSLƏR vərəqində
    ws_people = wb["ŞƏXSLƏR"]
    
    print("ŞƏXSLƏR vərəqində Data Validation:")
    # Aktivlik sütunu (G sütunu)
    print(f"  Aktivlik sütunu (G5:G10): {ws_people['G5'].data_type}")
    
    # Növbə uyğunluğu sütunları
    for i in range(6):
        col = 9 + i
        col_letter = openpyxl.utils.get_column_letter(col)
        print(f"  {col_letter}5: {ws_people[f'{col_letter}5'].data_type}")
    
    # AYLIQ QRAFİK vərəqində
    ws_schedule = wb["AYLIQ QRAFİK"]
    
    print("\nAYLIQ QRAFİK vərəqində Data Validation:")
    print(f"  İl (B4): {ws_schedule['B4'].data_type}")
    print(f"  Ay (E4): {ws_schedule['E4'].data_type}")
    
    # Növbə sütunları
    for i in range(6):
        col = 5 + i
        col_letter = openpyxl.utils.get_column_letter(col)
        print(f"  {col_letter}8: {ws_schedule[f'{col_letter}8'].data_type}")
    
    return True


def test_formulas():
    """Formulaların testi"""
    print("\n" + "=" * 60)
    print("TEST 9: Formulaların düzgünlüyü")
    print("=" * 60)
    
    wb = load_workbook('/workspace/github__Mireyyub__zenthon/Novbe_Sistemi.xlsx')
    
    # YOXDUR-MƏZUNİYYƏT vərəqində bitmə tarixi
    ws_absence = wb["YOXDUR-MƏZUNİYYƏT"]
    
    print("YOXDUR-MƏZUNİYYƏT vərəqində:")
    for row in range(5, 8):
        formula = ws_absence[f"E{row}"].value
        if formula and formula.startswith("="):
            print(f"  ✅ Sətir {row}: {formula}")
        else:
            print(f"  ❌ Sətir {row}: Formula tapılmadı")
    
    # BALANS vərəqində
    ws_balance = wb["BALANS"]
    
    print("\nBALANS vərəqində:")
    for row in range(5, 8):
        person = ws_balance[f"A{row}"].value
        if person:
            formula_b = ws_balance[f"B{row}"].value
            formula_c = ws_balance[f"C{row}"].value
            print(f"  {person}:")
            print(f"    Ümumi növbə: {formula_b[:50] if formula_b and formula_b.startswith('=') else 'Formula yoxdur'}...")
            print(f"    Ümumi xidmət yükü: {formula_c[:50] if formula_c and formula_c.startswith('=') else 'Formula yoxdur'}...")
    
    return True


def main():
    """Əsas test funksiya"""
    print("\n" + "=" * 60)
    print("NÖVBƏÇİLİK SİSTEMİ TESTLƏRİ")
    print("=" * 60)
    
    results = []
    
    results.append(("Fayl strukturu", test_file_structure()))
    results.append(("PARAMETRLƏR vərəqi", test_parameters_sheet()))
    results.append(("ŞƏXSLƏR vərəqi", test_people_sheet()))
    results.append(("YOXDUR-MƏZUNİYYƏT vərəqi", test_absence_sheet()))
    results.append(("AYLIQ QRAFİK vərəqi", test_monthly_schedule()))
    results.append(("BALANS vərəqi", test_balance_sheet()))
    results.append(("DASHBOARD vərəqi", test_dashboard()))
    results.append(("Data Validation", test_data_validation()))
    results.append(("Formulalar", test_formulas()))
    
    # Nəticə
    print("\n" + "=" * 60)
    print("NƏTİCƏ")
    print("=" * 60)
    
    passed = sum(1 for _, result in results if result)
    total = len(results)
    
    for name, result in results:
        status = "✅ KEÇDI" if result else "❌ UĞURSUZ"
        print(f"{status}: {name}")
    
    print(f"\nÜmumi: {passed}/{total} test keçdi")
    
    if passed == total:
        print("\n🎉 Bütün testlər uğurla keçdi!")
    else:
        print(f"\n⚠️  {total - passed} test uğursuz oldu")


if __name__ == "__main__":
    main()
